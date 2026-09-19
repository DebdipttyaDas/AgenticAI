"""Multi-format document parser supporting PDF, DOCX, TXT, and MD."""

import os
from pathlib import Path
from typing import Dict, Any, Tuple


class DocumentParser:
    """Extracts raw text and metadata from resumes and job descriptions across multiple formats."""

    @staticmethod
    def parse_file(file_path: str | Path) -> Tuple[str, Dict[str, Any]]:
        """Parse file based on extension and return (raw_text, metadata)."""
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

        suffix = path.suffix.lower()
        metadata = {
            "filename": path.name,
            "filepath": str(path.resolve()),
            "extension": suffix,
            "size_bytes": path.stat().st_size,
        }

        if suffix in [".txt", ".md"]:
            raw_text = path.read_text(encoding="utf-8", errors="replace")
        elif suffix == ".pdf":
            raw_text = DocumentParser._parse_pdf(path)
        elif suffix in [".docx", ".doc"]:
            raw_text = DocumentParser._parse_docx(path)
        else:
            # Fallback to UTF-8 text read
            raw_text = path.read_text(encoding="utf-8", errors="replace")

        return raw_text.strip(), metadata

    @staticmethod
    def _parse_pdf(path: Path) -> str:
        """Extract text using pypdf."""
        try:
            import pypdf
            reader = pypdf.PdfReader(str(path))
            pages = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                pages.append(text)
            return "\n\n".join(pages)
        except Exception as e:
            raise RuntimeError(f"Failed to parse PDF {path}: {str(e)}")

    @staticmethod
    def _parse_docx(path: Path) -> str:
        """Extract text using python-docx."""
        try:
            import docx
            doc = docx.Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            return "\n\n".join(paragraphs)
        except Exception as e:
            raise RuntimeError(f"Failed to parse DOCX {path}: {str(e)}")
